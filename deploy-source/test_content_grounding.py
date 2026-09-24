from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace

from docx import Document

from app.services.source_analyzer.parsers import parse_docx
from app.services.source_analyzer.knowledge_map import build_local_knowledge_map


with TemporaryDirectory() as td:
    path=Path(td)/"lesson.docx"
    doc=Document()
    doc.add_heading("BÀI HỌC", level=1)
    doc.add_paragraph("Mở đầu của giáo án")
    table=doc.add_table(rows=2, cols=2)
    table.cell(0,0).text="Hoạt động của cô"
    table.cell(0,1).text="Hoạt động của trẻ"
    table.cell(1,0).text="Cô tổ chức quan sát và đặt câu hỏi"
    table.cell(1,1).text="Trẻ quan sát, trả lời và nêu nhận xét"
    doc.add_paragraph("Kết luận sau hoạt động")
    doc.save(path)
    parsed=parse_docx(path)
    kinds=[u.metadata.get("kind") for u in parsed.units]
    assert kinds==["paragraph","paragraph","table","paragraph"], kinds
    assert parsed.metadata.get("preserved_document_order") is True

chunks=[]
for i in range(20):
    chunks.append(SimpleNamespace(id=f"c{i}", heading=f"Mục {i}", content=("Nội dung giáo án quan trọng "+str(i)+" ")*40))
km=build_local_knowledge_map("Bài kiểm thử", chunks)
assert len(km.topics)==20, len(km.topics)
assert all(t.source_chunk_ids for t in km.topics)
assert max(len(f.text) for t in km.topics for f in t.key_facts)>280
print("source-grounding-regression-ok")


from app.services.pedagogy_planner.service import _repair_plan_source_coverage
from app.services.pedagogy_planner.types import PedagogyPlan, PlannedSection

plan=PedagogyPlan(
    plan_title="Coverage test",
    lesson_strategy="source first",
    total_duration_seconds=1800,
    objectives=[],
    sections=[
        PlannedSection(section_key="explore_1",section_order=4,title="Khám phá 1",purpose="p",planned_duration_seconds=300,slide_count_hint=1),
        PlannedSection(section_key="explore_2",section_order=5,title="Khám phá 2",purpose="p",planned_duration_seconds=300,slide_count_hint=1),
        PlannedSection(section_key="explore_3",section_order=6,title="Khám phá 3",purpose="p",planned_duration_seconds=300,slide_count_hint=1),
    ],
)
fake_chunks=[SimpleNamespace(id=f"chunk-{i}",heading=f"Mục {i//2}") for i in range(9)]
plan=_repair_plan_source_coverage(plan,fake_chunks)
covered={cid for section in plan.sections for cid in section.source_chunk_ids}
assert covered=={f"chunk-{i}" for i in range(9)}, covered
assert all(section.slide_count_hint>=1 for section in plan.sections)
assert any("Độ phủ nguồn" in w for w in plan.warnings)
print("planner-source-coverage-regression-ok")


from app.services.lesson_writer.service import _section_ref_coverage

fake_section=SimpleNamespace(
    slides=[
        SimpleNamespace(
            source_refs=[SimpleNamespace(source_chunk_id="a")],
            interaction=None,
        )
    ]
)
coverage=_section_ref_coverage(fake_section,{"source_chunk_ids":["a","b","c","d"]})
assert abs(coverage-0.25)<1e-9, coverage
print("writer-source-coverage-regression-ok")


from app.services.lesson_writer.service import _draft_source_coverage

draft=SimpleNamespace(slides=[
    SimpleNamespace(source_refs=[SimpleNamespace(source_chunk_id="a"),SimpleNamespace(source_chunk_id="b")], interaction=None),
    SimpleNamespace(source_refs=[SimpleNamespace(source_chunk_id="c")], interaction=SimpleNamespace(source_chunk_ids=["d"])),
])
coverage,used,total=_draft_source_coverage(draft,{"sections":[
    {"section_key":"explore_1","source_chunk_ids":["a","b"]},
    {"section_key":"explore_2","source_chunk_ids":["c"]},
    {"section_key":"explore_3","source_chunk_ids":["d","e"]},
]})
assert (used,total)==(4,5)
assert abs(coverage-0.8)<1e-9, coverage
print("final-draft-source-coverage-regression-ok")


# Regression: V2 must enrich interactions before section validation inside _build_section.
from pathlib import Path as _Path
_service_text=_Path("app/services/lesson_writer/service.py").read_text()
_build_start=_service_text.index("def _build_section(")
_build_end=_service_text.find("\ndef ", _build_start+1)
_build_body=_service_text[_build_start:_build_end if _build_end>0 else None]
_ai_enrich=_build_body.index("candidate = enrich_section_v2(")
_ai_validate=_build_body.index("candidate = validate_and_repair_section(")
_local_enrich=_build_body.index("local = enrich_section_v2(")
_local_validate=_build_body.index("local = validate_and_repair_section(")
assert _ai_enrich < _ai_validate, "AI writer validation occurs before V2 interaction enrichment"
assert _local_enrich < _local_validate, "Local fallback validation occurs before V2 interaction enrichment"
assert "interaction_2" not in _build_body or _ai_enrich < _ai_validate
print("writer-enrich-before-validate-regression-ok")


import importlib.util as _importlib_util
_v2_spec=_importlib_util.spec_from_file_location("elearning_v2_test", _Path("../deploy-source/elearning_v2.py"))
assert _v2_spec and _v2_spec.loader
_v2_module=_importlib_util.module_from_spec(_v2_spec)
_v2_spec.loader.exec_module(_v2_module)
_design_profile=_v2_module._design_profile

_sections=[
    "introduction","objectives","warmup","lead_in",
    "explore_1","interaction_1","explore_2","interaction_2",
    "explore_3","interaction_3","practice","application","summary","closing",
]
_profiles=[_design_profile(key,0,"",key,[]) for key in _sections]
_layouts={p["layout"] for p in _profiles}
_visuals={p["visual"] for p in _profiles}
assert len(_layouts)>=12, _layouts
assert len(_visuals)>=9, _visuals
_explore_layouts={_design_profile("explore_1",i,"","Khám phá",[])["layout"] for i in range(3)}
assert len(_explore_layouts)==3, _explore_layouts
_practice_layouts={_design_profile("practice",i,"","Luyện tập",[])["layout"] for i in range(3)}
assert len(_practice_layouts)==3, _practice_layouts
print("slide-layout-diversity-regression-ok",len(_layouts),len(_visuals))


_map_profiles=[
    _design_profile("explore_1",i,"map","Kinh tuyến và vĩ tuyến",["Xác định vị trí trên bản đồ"])
    for i in range(6)
]
_map_layouts=[p["layout"] for p in _map_profiles]
assert _map_layouts.count("map-focus") <= 2, _map_layouts
_non_map=_design_profile("explore_1",0,"map","Vòng tuần hoàn của nước",["Nước bay hơi và ngưng tụ"])
assert _non_map["layout"]!="map-focus", _non_map
print("semantic-map-diversity-regression-ok",_map_layouts)


_interaction_map=_design_profile("interaction_1",0,"map","Xác định vị trí trên bản đồ",["Quan sát bản đồ và chọn đáp án"])
assert _interaction_map["layout"]=="quiz-card", _interaction_map
_practice_map=_design_profile("practice",1,"map","Luyện tập bản đồ",["Ghép vị trí với mô tả"])
assert _practice_map["layout"]=="matching-workspace", _practice_map
_application_map=_design_profile("application",0,"map","Vận dụng bản đồ",["Giải quyết tình huống thực tế"])
assert _application_map["layout"]=="real-world-scenario", _application_map
_geo_family=[
    _design_profile("explore_1",0,"map","Kinh tuyến và vĩ tuyến",["Xác định vị trí trên bản đồ"]),
    _design_profile("explore_1",1,"map","Kinh tuyến và vĩ tuyến",["Xác định vị trí trên bản đồ"]),
    _design_profile("explore_2",0,"map","Kinh tuyến và vĩ tuyến",["Xác định vị trí trên bản đồ"]),
    _design_profile("explore_3",0,"map","Kinh tuyến và vĩ tuyến",["Xác định vị trí trên bản đồ"]),
]
assert len({x["layout"] for x in _geo_family})>=3, _geo_family
print("interaction-role-layout-precedence-ok", [x["layout"] for x in _geo_family])


from app.services.visual_catalog import LAYOUTS as _VISUAL_LAYOUTS

_required_v2_layouts = {
    "cinematic-hero","milestone-checklist","question-spotlight","scenario-stage",
    "quiz-card","fill-blank-focus","evidence-choice","real-world-scenario",
    "concept-map","takeaway-cards","source-list","activity-board",
    "matching-workspace","sequence-workspace","comparison-2-column",
    "evidence-board","split-visual-explain","chart-focus","data-table",
    "process-timeline","cause-effect","map-focus","annotated-visual",
    "zoom-detail","full-bleed-annotated","visual-left-text-right",
    "text-left-visual-right","center-focus",
}
_catalog_by_key={item["key"]:item for item in _VISUAL_LAYOUTS}
_missing_layouts=sorted(_required_v2_layouts-set(_catalog_by_key))
assert not _missing_layouts, f"V2 layout hints missing from visual catalog: {_missing_layouts}"
_geometry_signatures={
    tuple(sorted((name,tuple(box)) for name,box in _catalog_by_key[key]["elements"].items()))
    for key in _required_v2_layouts
}
assert len(_geometry_signatures)>=24, f"V2 catalog geometry is too repetitive: {len(_geometry_signatures)} unique layouts"
assert _catalog_by_key["map-focus"]["elements"]["visual"][2] >= 60
assert _catalog_by_key["comparison-2-column"]["elements"]["body"] != _catalog_by_key["map-focus"]["elements"]["body"]
assert _catalog_by_key["process-timeline"]["elements"]["visual"] != _catalog_by_key["annotated-visual"]["elements"]["visual"]
print("v2-visual-catalog-regression-ok",len(_required_v2_layouts),"layouts",len(_geometry_signatures),"geometries")


# Semantic slide-layout regression: layouts must follow the learning task,
# not merely the subject vocabulary.
def _layout(title, bullets=None, visual="", index=0, previous=None, interaction_type=""):
    return _design_profile(
        "explore_1", index, visual, title, bullets or [],
        interaction_type=interaction_type,
        previous_layouts=previous or [],
    )["layout"]

assert _layout("Phân tích bảng số liệu dân số", ["Bảng số liệu cho biết..."]) == "data-table"
assert _layout("Nhận xét biểu đồ cơ cấu", ["Quan sát biểu đồ tròn..."]) == "chart-focus"
assert _layout("Nguyên nhân và hệ quả của hiện tượng", ["Nguyên nhân...", "Hệ quả..."]) == "cause-effect"
assert _layout("Quy trình hình thành", ["Các bước diễn ra theo trình tự..."]) == "process-timeline"
assert _layout("So sánh hai khu vực", ["Khác nhau về..."]) == "comparison-2-column"

# Concept vocabulary alone must not turn every geography slide into a map.
coord_layout = _layout(
    "Khái niệm kinh tuyến và vĩ tuyến",
    ["Kinh tuyến và vĩ tuyến được dùng để xác định tọa độ địa lí."],
    visual="map",
)
assert coord_layout != "map-focus", coord_layout

map_layout = _layout(
    "Xác định vị trí trên bản đồ",
    ["Quan sát bản đồ và xác định vị trí địa lí của khu vực."],
    visual="map",
)
assert map_layout == "map-focus", map_layout

map_layout_2 = _layout(
    "Khai thác bản đồ",
    ["Quan sát bản đồ và nhận xét sự phân bố."],
    visual="map",
    previous=["map-focus"],
)
assert map_layout_2 != "map-focus", map_layout_2

practice_match = _design_profile(
    "practice", 0, "quiz_ui", "Luyện tập", ["Ghép nội dung"],
    interaction_type="matching",
    previous_layouts=[],
)["layout"]
assert practice_match == "matching-workspace", practice_match
print("semantic-layout-selection-regression-ok")
